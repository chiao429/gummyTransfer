"""Fixed-scope report versions from stored data; regeneration never fetches suppliers."""

import hashlib
import json
import math
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, insert, select

from gummy.analysis import AnalysisQuery, analysis_data
from gummy.db import operation_reports, report_versions, tasks, utc_now
from gummy.monitor import active_plan_overview, report_content_verified

SCHEMA = 4


def checksum(value):
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()
    ).hexdigest()


def fixed_report_data(db, scope):
    query = AnalysisQuery.model_validate(scope)
    data = analysis_data(db, query)
    if (
        query.snapshot is not None
        and query.snapshot != "0"
        and set(query.snapshot.split(",")) != {str(r["id"]) for r in data["runs"]}
    ):
        raise ValueError("固定批次已不存在或來源不符。")
    digest = checksum(
        {
            "scope": data["query"].model_dump(mode="json"),
            "rows": data["rows"],
            "runs": data["runs"],
            "summary": data["summary"],
            "groups": data["groups"],
            "source_evidence": {str(r["id"]): db.evidence(r["id"]) for r in data["runs"]},
        }
    )
    return data, digest


def prepare_report(db, query):
    data, digest = fixed_report_data(db, query.model_dump(mode="json"))
    if not data["rows"]:
        raise ValueError("這個範圍沒有商品，請調整查詢。")
    from gummy.operations import TaskInput

    return TaskInput(
        kind="report",
        source=query.source,
        run_id=max(r["id"] for r in data["runs"]),
        report_scope=data["query"].model_dump(mode="json"),
        report_digest=digest,
    ), data


def register_version(conn, spec, task_id):
    scope_hash = checksum(spec.report_scope)
    parent = (
        conn.execute(select(report_versions).where(report_versions.c.id == spec.parent_report))
        .mappings()
        .first()
        if spec.parent_report
        else None
    )
    if spec.parent_report and (
        not parent
        or parent["source"] != spec.source
        or parent["scope"] != spec.report_scope
        or parent["data_digest"] != spec.report_digest
    ):
        raise ValueError("重產版本範圍不同。")
    version = (
        conn.execute(
            select(func.max(report_versions.c.version)).where(
                report_versions.c.source == spec.source, report_versions.c.scope_hash == scope_hash
            )
        ).scalar()
        or 0
    ) + 1
    conn.execute(
        insert(report_versions).values(
            id=task_id,
            source=spec.source,
            scope=spec.report_scope,
            scope_hash=scope_hash,
            data_digest=spec.report_digest,
            parent_id=spec.parent_report,
            version=version,
            created_at=utc_now(),
        )
    )


def version(db, report_id, source):
    with db.engine.connect() as conn:
        row = (
            conn.execute(
                select(report_versions).where(
                    report_versions.c.id == report_id, report_versions.c.source == source
                )
            )
            .mappings()
            .first()
        )
    if not row:
        raise ValueError("找不到此來源的報表版本。")
    return dict(row)


def tables_for(data):
    q = data["query"]
    summary = data["summary"]
    scope = json.dumps(q.model_dump(mode="json"), ensure_ascii=False, sort_keys=True)
    groups = [
        [
            "分類／品牌",
            "價格帶",
            "商品數",
            "有價格商品數",
            "平均價格 USD",
            "有評價數商品數",
            "已知累計評價數合計",
            "留存評論樣本數",
            "有來源證據商品數",
        ]
    ]
    groups += [
        [
            g["group_value"],
            g["price_band_label"] or "",
            g["product_count"],
            g["priced_product_count"],
            float(g["average_price"]) if g["average_price"] is not None else "",
            g["rating_count_product_count"],
            g["known_rating_count_sum"],
            g["retained_review_sample_count"],
            g["evidence_product_count"],
        ]
        for g in data["groups"]
    ]
    products = [
        [
            "批次",
            "ASIN",
            "分類",
            "商品名稱",
            "品牌",
            "價格 USD",
            "評分",
            "評價數",
            "採集時間 UTC",
            "來源證據",
            "評論樣本數",
        ]
    ]
    products += [
        [
            r["run_id"],
            r["asin"],
            r["category"],
            r["title"],
            r["brand"],
            float(r["price"]) if r["price"] is not None else "",
            float(r["rating"]) if r["rating"] is not None else "",
            r["rating_count"],
            r["observed_at"],
            r["evidence_status"],
            r["review_sample_count"],
        ]
        for r in data["rows"]
    ]
    lineage = [["批次", "搜尋分類", "關鍵字", "市場", "完成時間 UTC", "此報表採用商品數"]]
    lineage += [
        [
            r["id"],
            r["parameters"]["category"],
            r["parameters"]["query"],
            r["parameters"]["country"],
            r["completed_at"],
            sum(row["run_id"] == r["id"] for row in data["rows"]),
        ]
        for r in data["runs"]
    ]
    return {
        "報表說明": [
            ["項目", "內容"],
            ["報表資料版本", SCHEMA],
            ["來源", q.source],
            ["台灣開始日期", str(q.date_from or "最早")],
            ["台灣結束日期", str(q.date_to or "截至選定批次")],
            ["固定成功批次", q.snapshot],
            ["網站核對截止版本", q.review_snapshot],
            ["固定查詢條件", scope],
            ["取樣商品數", summary["product_count"]],
            ["有價格商品數", summary["priced_product_count"]],
            [
                "取樣平均價格 USD",
                float(summary["average_price"])
                if summary["average_price"] is not None
                else "未提供",
            ],
            ["已知累計評價數合計", summary["known_rating_count_sum"]],
            ["評價數定義", "選定期間每個國家／ASIN 最新觀測的累計數；不是本期新增評論數。"],
            ["資料範圍", "有限取樣，不是整體市場；真實月銷量及銷售金額未提供。"],
            ["選取順序", data["latest_selection_basis"]],
            ["匯出 API 用量", 0],
        ],
        "取樣統計": groups,
        "商品明細": products,
        "批次來源": lineage,
    }


def export_fixed_report(db, spec, destination):
    data, digest = fixed_report_data(db, spec.report_scope)
    if digest != spec.report_digest:
        raise ValueError("固定資料已改變，請重新查詢與預覽。")
    tables = tables_for(data)
    book = Workbook()
    book.remove(book.active)
    edge = Side(style="thin", color="B8C5D1")
    for name, rows in tables.items():
        sheet = book.create_sheet(name)
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        widths = [70 if h in {"內容", "商品名稱"} else 28 for h in rows[0]]
        for i, row in enumerate(rows, 1):
            for j, value in enumerate(row, 1):
                c = sheet.cell(i, j, value)
                if isinstance(value, str):
                    c.data_type = "s"
                c.font = Font(name="Arial", size=11, color="26384A")
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                if isinstance(value, float):
                    c.number_format = "0.0" if "評分" in rows[0][j - 1] else "#,##0.00"
                elif isinstance(value, int):
                    c.number_format = "#,##0"
                if i == 1:
                    c.font = Font(name="Arial", bold=True, color="FFFFFF")
                    c.fill = PatternFill("solid", fgColor="18324D")
            sheet.row_dimensions[i].height = (
                42
                if i == 1
                else min(
                    409,
                    max(
                        30,
                        max(
                            sum(
                                max(
                                    1,
                                    math.ceil(
                                        sum(
                                            2
                                            if unicodedata.east_asian_width(ch) in {"W", "F"}
                                            else 1
                                            for ch in line
                                        )
                                        / (widths[j] - 3)
                                    ),
                                )
                                for line in str(value if value is not None else "").split("\n")
                            )
                            * 17
                            + 8
                            for j, value in enumerate(row)
                        ),
                    ),
                )
            )
        for j, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(j)].width = width
        if name != "報表說明":
            sheet.auto_filter.ref = sheet.dimensions
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("xb") as stream:
        destination.chmod(0o600)
        book.save(stream)
    actual = load_workbook(destination, read_only=True, data_only=False)
    try:
        if actual.sheetnames != list(tables):
            raise ValueError("報表分頁不符。")
        for name, rows in tables.items():
            observed = list(actual[name].iter_rows(values_only=True))
            if len(observed) != len(rows):
                raise ValueError("報表資料列不符。")
            for got, want in zip(observed, rows):
                if list(got) != [None if v == "" else v for v in want]:
                    raise ValueError("報表讀回核對失敗。")
            if any(c.data_type == "f" for row in actual[name] for c in row):
                raise ValueError("報表不允許公式。")
    finally:
        actual.close()
    sha = hashlib.sha256(destination.read_bytes()).hexdigest()
    record = {
        "schema_version": 1,
        "verification_status": "PASS",
        "source": spec.source,
        "through_run_id": spec.run_id,
        "sha256": sha,
        "workbook_schema_version": SCHEMA,
        "scope_hash": checksum(spec.report_scope),
        "data_digest": digest,
        "export_engine": "python-openpyxl",
    }
    side = Path(str(destination) + ".verified.json")
    with side.open("x") as stream:
        side.chmod(0o600)
        json.dump(record, stream, ensure_ascii=False, default=str)
    return record


def center_rows(db, settings, source):
    with db.engine.connect() as conn:
        versions = [
            dict(r)
            for r in conn.execute(
                select(report_versions, tasks.c.status, tasks.c.run_id)
                .join(tasks, tasks.c.id == report_versions.c.id)
                .where(report_versions.c.source == source)
                .order_by(report_versions.c.created_at.desc(), report_versions.c.id)
                .limit(100)
            ).mappings()
        ]
        stored = {
            r["task_id"]: dict(r)
            for r in conn.execute(
                select(operation_reports).where(operation_reports.c.source == source)
            ).mappings()
        }
        fixed_ids = set(
            conn.execute(
                select(report_versions.c.id).where(report_versions.c.source == source)
            ).scalars()
        )
    for row in versions:
        report = stored.get(row["id"])
        row["verified"] = bool(
            report and report_content_verified(Path(report["path"]), source, row["run_id"])
        )
        row["path"] = report["path"] if report else None
    plan = active_plan_overview(settings, db)
    daily = [b for b in plan.get("batches", []) if plan.get("source") == source and b.get("report")]
    legacy = [r for key, r in stored.items() if key not in fixed_ids]
    return versions, daily, legacy
