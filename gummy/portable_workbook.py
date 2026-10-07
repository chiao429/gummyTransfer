"""Python-only portfolio XLSX export and source verification; no network access."""

import hashlib
import json
import math
import re
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gummy.reference import reference_coverage
from gummy.review_status import MAX_MEMBER_BYTES, MAX_WORKBOOK_BYTES, MAX_ZIP_ENTRIES
from gummy.workbook_data import attach_web_reviews, portfolio_payload

PRODUCT_HEADERS = [
    "批次",
    "搜尋分類（待人工確認）",
    "ASIN",
    "商品標題",
    "品牌",
    "價格 USD",
    "幣別",
    "評分",
    "評價數",
    "材質",
    "結構化制式",
    "標題候選（未驗證）",
    "尺寸",
    "顏色",
    "來源證據",
    "批次對帳",
    "商品差異數",
    "擷取時間 UTC",
]
AUDIT_HEADERS = [
    "批次",
    "搜尋分類",
    "ASIN",
    "證據狀態",
    "批次對帳",
    "商品差異數",
    "來源標題",
    "入庫標題",
    "來源價格",
    "入庫價格 USD",
    "人工分類確認",
    "人工來源核對",
    "人工備註",
    "擷取時間 UTC",
    "Amazon 網址（複製開啟）",
]
SOURCE_HEADERS = [
    "批次",
    "搜尋分類",
    "ASIN",
    "來源證據狀態",
    "來源標題",
    "入庫標題",
    "來源價格",
    "入庫原價格",
    "入庫價格 USD",
    "來源品牌",
    "入庫品牌",
    "來源材質",
    "入庫材質",
    "來源制式",
    "入庫制式",
    "標題候選（未驗證）",
    "來源尺寸",
    "入庫尺寸",
    "來源顏色",
    "入庫顏色",
    "來源評分",
    "入庫評分",
    "來源評價數",
    "入庫評價數",
    "擷取時間 UTC",
]
SECTIONS = {
    "search_listing": "搜尋結果",
    "product_summary": "商品明細摘要",
    "product_information": "商品資訊",
    "product_details": "商品詳情",
    "top_review_sample": "精選評論樣本",
}


def text(value):
    return None if value is None or value == "" else str(value)


def number(value):
    if value is None or value == "":
        return None
    result = float(value)
    if not math.isfinite(result):
        raise ValueError("Non-finite workbook source number")
    return result


def timestamp(value):
    if not value:
        return None
    return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(UTC).replace(tzinfo=None)


def evidence_label(row):
    return (
        "已留存選定欄位"
        if row["evidence_status"] == "captured_selected_fields"
        else "本批未留存來源欄位"
    )


def product_values(row):
    return [
        number(row["run_id"]),
        *[text(row[k]) for k in ("category", "asin", "title", "brand")],
        number(row["price"]),
        text(row["currency"]),
        number(row["rating"]),
        number(row["rating_count"]),
        *[
            text(row[k])
            for k in (
                "material",
                "measurement_system",
                "title_measurement_candidates",
                "size",
                "color",
            )
        ],
        evidence_label(row),
        text(row["run_verification"]),
        number(row["product_issue_count"]),
        timestamp(row["observed_at"]),
    ]


def source_values(row):
    fields = row["source_audit"]
    keys = [
        "source_title",
        "stored_title",
        "source_price",
        "stored_price_raw",
        "stored_price_usd",
        "source_brand",
        "stored_brand",
        "source_material",
        "stored_material",
        "source_measurement_system",
        "stored_measurement_system",
        "title_measurement_candidates_unverified",
        "source_size",
        "stored_size",
        "source_color",
        "stored_color",
        "source_rating",
        "stored_rating",
        "source_rating_count",
        "stored_rating_count",
    ]
    numeric = {"stored_price_usd", "stored_rating", "stored_rating_count"}
    return [
        number(row["run_id"]),
        text(row["category"]),
        text(row["asin"]),
        evidence_label(row),
        *[(number if key in numeric else text)(fields[key]) for key in keys],
        timestamp(fields["observed_at"]),
    ]


def payload_from_db(db, source, run_id, *, version=3, watermark=None, web=True):
    if version not in (1, 2, 3):
        raise ValueError("Unsupported workbook schema version")
    snapshots = [
        (*db.report(source, item), db.evidence(item))
        for item in db.successful_run_ids(source, run_id)
    ]
    payload = portfolio_payload(
        snapshots, reference_coverage(db.category_stats(source, through_run_id=run_id), source)
    )
    if payload["through_run_id"] != run_id:
        raise ValueError("Workbook cutoff is not a successful batch")
    payload["workbook_schema_version"] = version
    if version == 2:
        from gummy.quality import sample_label_rows

        payload["review_samples"] = sample_label_rows(payload["review_samples"], None, "")
    if web:
        attach_web_reviews(db, payload, watermark)
    return payload


def workbook_tables(payload):
    """Retain the existing sheet names, column order and manual K–M positions."""
    tables = []

    def add(name, title, headers, rows, metadata, header=3):
        tables.append(
            dict(
                name=name, title=title, headers=headers, rows=rows, metadata=metadata, header=header
            )
        )

    categories = []
    for outside, items in [
        (False, payload["coverage"]["categories"]),
        (True, payload["coverage"]["outside_reference"]),
    ]:
        for index, item in enumerate(items, 1):
            short = re.sub(r"[\\/:*?\[\]]", "-", item["category"].split("(")[0]).strip()
            name = (("其他" if outside else "分類") + f"{index:02d}-" + short)[:31]
            categories.append(
                {**item, "type": "參考表以外" if outside else "原表分類", "tab": name}
            )
    sampled = [item for item in categories if item["run_count"]]
    add(
        "分類總覽",
        "跨批次分類總覽｜成功快照取樣",
        ["類型", "原表 Category", "狀態", "不重複 ASIN", "成功批次", "最新 run_id", "分類分頁"],
        [
            [
                c["type"],
                text(c["category"]),
                ("模擬取樣" if payload["source"] == "demo" else "已取樣")
                if c["run_count"]
                else "未取樣",
                number(c["distinct_product_count"]),
                number(c["run_count"]),
                number(c["latest_run_id"]),
                c["tab"] if c["run_count"] else None,
            ]
            for c in categories
        ],
        [
            "來源",
            payload["source"],
            "截至批次",
            payload["through_run_id"],
            "成功批次",
            payload["run_count"],
            f"分類取樣 {payload['coverage']['sampled_count']}/"
            f"{payload['coverage']['reference_total']}",
        ],
        header=4,
    )
    add(
        "規格組合",
        "跨分類規格組合｜成功快照取樣",
        [
            "搜尋分類",
            "材質",
            "公制／英制",
            "尺寸",
            "顏色",
            "觀測列數",
            "不重複 ASIN",
            "有價格列數",
            "平均價 USD",
            "標題候選列數",
            "來源留存列數",
            "來源對帳 PASS 列數",
            "ASIN",
            "批次:ASIN",
        ],
        [
            [
                *[
                    text(g[k])
                    for k in ("category", "material", "measurement_system", "size", "color")
                ],
                *[
                    number(g[k])
                    for k in (
                        "observation_count",
                        "distinct_asin_count",
                        "priced_observation_count",
                        "average_price_usd",
                        "title_candidate_observation_count",
                        "source_evidence_observation_count",
                        "source_pass_observation_count",
                    )
                ],
                "; ".join(g["asins"]),
                "; ".join(g["run_asins"]),
            ]
            for g in payload["combinations"]
        ],
        ["同一 ASIN 可出現在多組；標題候選未驗證；來源對帳 PASS 不是人工確認。"],
    )
    observations = payload["observations"]
    add(
        "跨分類商品",
        "跨分類商品｜每列是一批次中的一件商品",
        PRODUCT_HEADERS,
        [product_values(row) for row in observations],
        [
            "觀測列",
            payload["observation_count"],
            "不重複 ASIN",
            payload["distinct_asin_count"],
            "同一 ASIN 在不同批次保留各自的價格與來源證據。",
        ],
    )
    add(
        "對帳索引",
        "跨批次來源對帳索引｜需人工核對來源內容",
        AUDIT_HEADERS,
        [
            [
                number(row["run_id"]),
                text(row["category"]),
                text(row["asin"]),
                evidence_label(row),
                text(row["run_verification"]),
                number(row["product_issue_count"]),
                text(row["source_title"]),
                text(row["title"]),
                text(row["source_price"]),
                number(row["price"]),
                "待確認",
                "待確認",
                None,
                timestamp(row["observed_at"]),
                f"https://www.amazon.com/dp/{row['asin']}"
                if payload["source"] == "rapidapi" and re.fullmatch(r"[A-Z0-9]{10}", row["asin"])
                else None,
            ]
            for row in observations
        ],
        ["PASS 僅核對留存欄位；K–M 為保留的 Excel 人工欄，非評論判斷。"],
    )
    add(
        "來源欄位對照",
        "跨批次來源欄位對照｜僅留存的選定欄位",
        SOURCE_HEADERS,
        [source_values(row) for row in observations],
        ["來源欄空白可能是供應商未提供或舊批次未留存；不能把空白當成一致。"],
    )
    attrs = payload.get("additional_attributes", [])
    if attrs:
        add(
            "來源附加欄位",
            "供應商附加欄位與精選評論樣本",
            [
                "批次",
                "搜尋分類",
                "ASIN",
                "擷取時間 UTC",
                "來源區塊",
                "來源欄名",
                "來源原值",
                "分類狀態",
            ],
            [
                [
                    number(r["run_id"]),
                    text(r["category"]),
                    text(r["asin"]),
                    timestamp(r["observed_at"]),
                    SECTIONS[r["source_section"]],
                    text(r["field_name"]),
                    text(r["source_value"]),
                    "未分類；待人工確認",
                ]
                for r in attrs
            ],
            ["保留供應商原文。評論僅為精選樣本；sales_volume 是非精確購買提示，不是實際銷量。"],
        )
    version = payload["workbook_schema_version"]
    mappings = payload.get("specification_mappings", [])
    if version >= 2 and mappings:
        add(
            "規格映射",
            "來源規格與標準欄位對照",
            [
                "批次",
                "搜尋分類",
                "ASIN",
                "來源區塊",
                "來源欄名",
                "來源原值",
                "標準欄位",
                "候選值",
                "既有入庫值",
                "狀態",
                "映射版本",
            ],
            [
                [
                    number(r["run_id"]),
                    *[
                        text(r[k])
                        for k in (
                            "category",
                            "asin",
                            "source_section",
                            "field_name",
                            "source_value",
                            "dimension",
                            "candidate_value",
                            "stored_value",
                            "status",
                            "mapping_version",
                        )
                    ],
                ]
                for r in mappings
            ],
            ["欄名映射版本", payload["specification_mapping_version"], "候選未套用，衝突須核對。"],
        )
    reviews = payload.get("review_samples", [])
    if version >= 2 and reviews:
        rows = [
            [
                number(r["run_id"]),
                *[text(r[k]) for k in ("category", "asin", "review_id")],
                number(r["review_star_rating"]),
                *[text(r[k]) for k in ("review_date", "review_title", "review_comment")],
                *(
                    [
                        *[
                            text(r[k])
                            for k in (
                                "theme",
                                "sentiment",
                                "annotation_state",
                                "annotation_note",
                                "annotation_reviewer",
                            )
                        ],
                        number(r["annotation_revision"]),
                        timestamp(r["annotation_updated_at"]),
                        timestamp(r["observed_at"]),
                    ]
                    if version == 2
                    else [timestamp(r["observed_at"]), payload["source"]]
                ),
            ]
            for r in reviews
        ]
        headers = [
            "批次",
            "搜尋分類",
            "ASIN",
            "評論 ID",
            "星等",
            "評論日期原文",
            "評論標題原文",
            "評論全文原文",
            "擷取時間 UTC",
            "資料來源",
        ]
        add(
            "評論標註" if version == 2 else "評論來源",
            "採集到的評論來源資料",
            headers if version == 3 else [*headers[:8], *[""] * 8],
            rows,
            [
                "來源" if version == 3 else "網站儲存快照",
                payload["source"] if version == 3 else payload.get("web_review_snapshot", 0),
                "每列為留存的精選評論樣本，不代表全部評論。",
            ],
        )
    for category in sampled:
        rows = [row for row in observations if row["category"] == category["category"]]
        add(
            category["tab"],
            category["category"] + "｜搜尋分類待人工確認",
            PRODUCT_HEADERS,
            [product_values(row) for row in rows],
            [
                "分類觀測列",
                len(rows),
                "不重複 ASIN",
                category["distinct_product_count"],
                "成功批次",
                category["run_count"],
                "取樣資料，不代表市場總量。",
            ],
        )
    if "web_reviews" in payload:
        add(
            "網站核對",
            "網站核對結果",
            [
                "批次",
                "搜尋分類",
                "ASIN",
                "核對狀態",
                "分類判斷",
                "來源核對",
                "填寫人（自填）",
                "版本",
                "核對時間 UTC",
                "備註",
            ],
            [
                [
                    number(r["run_id"]),
                    *[
                        text(r[k])
                        for k in (
                            "category",
                            "asin",
                            "review_state",
                            "review_classification",
                            "review_source",
                            "reviewer",
                        )
                    ],
                    number(r["review_revision"]),
                    timestamp(r["review_updated_at"]),
                    text(r["review_note"]),
                ]
                for r in payload["web_reviews"]
            ],
            ["網站儲存快照", payload["web_review_snapshot"], "網站與 Excel 人工欄未互相同步。"],
        )
    names = [table["name"] for table in tables]
    if len(names) != len(set(names)):
        raise ValueError("Workbook sheet names collide")
    return tables


def read_book(path):
    """Bound XML before parsing untrusted editor output; reject formulas everywhere."""
    if path.stat().st_size > MAX_WORKBOOK_BYTES:
        raise ValueError("Workbook is too large")
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES or len({e.filename for e in entries}) != len(entries):
            raise ValueError("Workbook archive is invalid")
        if sum(e.file_size for e in entries) > 20 * MAX_MEMBER_BYTES:
            raise ValueError("Workbook expanded contents are too large")
        for entry in entries:
            if entry.file_size > MAX_MEMBER_BYTES:
                raise ValueError("Workbook member is too large")
            if entry.filename.endswith(".xml"):
                content = archive.read(entry)
                if b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
                    raise ValueError("Workbook XML declarations are not allowed")
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in book:
            if sheet.max_row is None or sheet.max_column is None:
                sheet.calculate_dimension(force=True)
            if sheet.max_row > 5000 or sheet.max_column > 40:
                raise ValueError("Workbook table bounds are too large")
            for row in sheet:
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("Workbook contains a formula")
    except Exception:
        book.close()
        raise
    return book


def _plain(value):
    if isinstance(value, str) and value.startswith("'") and re.match(r"\s*[=+\-@]", value[1:]):
        return value[1:]
    return None if value == "" else value


def _same(actual, expected):
    if actual == expected:
        return True
    actual = _plain(actual)
    if isinstance(expected, datetime):
        return isinstance(actual, datetime) and abs((actual - expected).total_seconds()) <= 1
    return actual == expected


def verify_book(payload, path):
    tables = workbook_tables(payload)
    book = read_book(Path(path))
    try:
        if book.sheetnames != [table["name"] for table in tables]:
            raise ValueError("Workbook sheet order or membership differs")
        for table in tables:
            sheet = book[table["name"]]
            for offset, row in enumerate(table["rows"], table["header"] + 1):
                for column, expected in enumerate(row, 1):
                    actual = sheet.cell(offset, column).value
                    if table["name"] == "對帳索引" and column in (11, 12, 13):
                        if actual is not None and (
                            not isinstance(actual, str) or len(actual) > 2000
                        ):
                            raise ValueError("Manual review must be bounded plain text")
                    elif not _same(actual, expected):
                        raise ValueError(
                            f"{table['name']}!{get_column_letter(column)}{offset} differs"
                        )
            end = table["header"] + len(table["rows"])
            for index, row in enumerate(
                sheet.iter_rows(min_row=table["header"] + 1), table["header"] + 1
            ):
                for column, cell in enumerate(row, 1):
                    if (index > end or column > len(table["headers"])) and cell.value not in (
                        None,
                        "",
                    ):
                        raise ValueError("Workbook has unexpected data outside the table")
        for sheet, cells in {
            "分類總覽": {
                "B2": payload["source"],
                "D2": payload["through_run_id"],
                "F2": payload["run_count"],
            },
            **(
                {"網站核對": {"B2": payload["web_review_snapshot"]}}
                if "web_reviews" in payload
                else {}
            ),
        }.items():
            for coordinate, value in cells.items():
                if book[sheet][coordinate].value != value:
                    raise ValueError("Workbook metadata differs from the snapshot")
    finally:
        book.close()
    return {
        "status": "PASS",
        "source": payload["source"],
        "through_run_id": payload["through_run_id"],
        "sheets_checked": len(tables),
        "observations_checked": len(payload["observations"]),
    }


def verify_from_db(db, source, run_id, path):
    record_path = Path(str(path) + ".verified.json")
    record = json.loads(record_path.read_text()) if record_path.is_file() else {}
    version = record.get("workbook_schema_version", 1 if record else 3)
    watermark = record.get("web_review_snapshot")
    if type(version) is not int or version not in (1, 2, 3):
        raise ValueError("Invalid workbook version")
    if watermark is not None and (type(watermark) is not int or not 0 <= watermark < 2**63):
        raise ValueError("Invalid web review snapshot")
    payload = payload_from_db(
        db, source, run_id, version=version, watermark=watermark, web=watermark is not None
    )
    return verify_book(payload, path)


def export_book(db, source, run_id, destination, review=None):
    destination = Path(destination)
    record_path = Path(str(destination) + ".verified.json")
    if destination.exists() or record_path.exists():
        raise FileExistsError("Report or verification record already exists")
    payload = payload_from_db(db, source, run_id)
    tables = workbook_tables(payload)
    if review:
        prior = read_book(Path(review))
        try:
            prior_id = prior["分類總覽"]["D2"].value
            if (
                type(prior_id) not in (int, float)
                or int(prior_id) != prior_id
                or not 0 < prior_id <= run_id
            ):
                raise ValueError("Review workbook cutoff is invalid")
            verify_from_db(db, source, int(prior_id), Path(review))
            manual = {
                (int(row[0]), row[2]): row[10:13]
                for row in prior["對帳索引"].iter_rows(min_row=4, values_only=True)
                if row[0] is not None
            }
        finally:
            prior.close()
        for table in tables:
            if table["name"] == "對帳索引":
                for row in table["rows"]:
                    if (int(row[0]), row[2]) in manual:
                        row[10:13] = list(manual[int(row[0]), row[2]])
    book = Workbook()
    book.remove(book.active)
    side = Side(style="thin", color="B8C5D1")
    for table in tables:
        sheet = book.create_sheet(table["name"])
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "D" + str(table["header"] + 1)
        for row_number, values in [
            (1, [table["title"]]),
            (2, table["metadata"]),
            (table["header"], table["headers"]),
            *[(index, row) for index, row in enumerate(table["rows"], table["header"] + 1)],
        ]:
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row_number, column, value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Source text never becomes an Excel formula.
                    cell.quotePrefix = bool(re.match(r"\s*[=+\-@]", value))
                cell.font = Font(name="Arial", size=11, color="26384A")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(value, float):
                    heading = table["headers"][column - 1]
                    if ("價格" in heading or "平均價" in heading) and "數" not in heading:
                        cell.number_format = "#,##0.00"
                    elif "評分" in heading or "星等" in heading:
                        cell.number_format = "0.0"
                    else:
                        cell.number_format = "#,##0"
                if row_number >= table["header"]:
                    cell.border = Border(left=side, right=side, top=side, bottom=side)
                    if row_number == table["header"]:
                        cell.fill = PatternFill("solid", fgColor="18324D")
                        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
                    elif table["name"] == "對帳索引" and column in (11, 12, 13):
                        cell.fill = PatternFill("solid", fgColor="FFF6DF")
            sheet.row_dimensions[row_number].height = (
                42
                if row_number == table["header"]
                else (
                    min(409, max(54, max((len(str(v or "")) // 42 + 1) * 15 for v in values)))
                    if row_number > table["header"]
                    else 30
                )
            )
        sheet["A1"].font = Font(name="Arial", size=14, bold=True, color="18324D")
        title_columns = min(8, len(table["headers"]))
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=title_columns)
        if len(table["metadata"]) == 1:
            sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=title_columns)
        for column, heading in enumerate(table["headers"], 1):
            width = (
                60
                if any(word in heading for word in ("標題", "原文", "原值", "備註", "網址"))
                else (38 if "分類" in heading or "Category" in heading else 23)
            )
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.auto_filter.ref = (
            f"A{table['header']}:{get_column_letter(len(table['headers']))}{sheet.max_row}"
        )
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("xb") as handle:
        destination.chmod(0o600)
        book.save(handle)
    result = verify_book(payload, destination)
    with destination.open("rb") as handle:
        digest = hashlib.file_digest(handle, "sha256").hexdigest()
    record = {
        "schema_version": 1,
        "source": source,
        "through_run_id": run_id,
        "verification_status": "PASS",
        "sha256": digest,
        "web_review_snapshot": payload["web_review_snapshot"],
        "workbook_schema_version": 3,
        "export_engine": "python-openpyxl",
    }
    with record_path.open("x", encoding="utf-8") as handle:
        record_path.chmod(0o600)
        json.dump(record, handle, ensure_ascii=False, sort_keys=True)
    return {**result, "output": str(destination.resolve())}
