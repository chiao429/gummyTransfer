"""Exact occurrence workbooks. Rebuilds use persisted batch IDs only."""

import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import insert, select

from gummy.db import operation_reports, tasks, utc_now
from gummy.recurrence import Rule, occurrence
from gummy.verification import verify_snapshot


def report_tables(db, occurrence_id):
    row = occurrence(db, occurrence_id)
    spec = Rule.model_validate(row["spec"])
    complete = len(row["results"]) == len(spec.categories)
    tables = {
        "執行說明": [
            ["項目", "內容"],
            ["排程", spec.name],
            ["執行識別", row["id"]],
            ["預定時間 UTC", row["scheduled_at"]],
            ["規則版本", row["revision"]],
            ["完整性", "完整" if complete else "不完整：部分分類未成功，沒有以舊資料補入"],
            ["成功分類數", len(row["results"])],
            ["預定分類數", len(spec.categories)],
            ["匯出 API 用量", 0],
            ["評價數定義", "來源累計評價數，不是本期新增評論數；不同分類可重複同一 ASIN"],
            ["範圍", "有限取樣；未提供真實銷售金額；不加入人工或 AI 評論判斷"],
        ],
        "分類結果": [["順序", "分類", "關鍵字", "狀態", "批次", "商品數"]],
        "商品明細": [
            [
                "批次",
                "分類",
                "ASIN",
                "商品名稱",
                "品牌",
                "價格",
                "幣別",
                "評分",
                "累計評價數",
                "觀測時間 UTC",
                "来源規格",
            ]
        ],
        "留存來源": [["批次", "ASIN", "來源欄位 JSON"]],
    }
    successes = {r["index"]: r for r in row["results"]}
    for index, category in enumerate(spec.categories):
        item = successes.get(index)
        if item is None:
            tables["分類結果"].append(
                [index + 1, category.category, category.query, "未成功／未執行", None, None]
            )
            continue
        run, products = db.report(spec.source, item["run_id"])
        evidence = db.evidence(run["id"])
        if verify_snapshot(run, products, evidence)["status"] != "PASS":
            raise ValueError("來源一致性核對未通過")
        tables["分類結果"].append(
            [index + 1, category.category, category.query, "成功", run["id"], len(products)]
        )
        for product in products:
            tables["商品明細"].append(
                [
                    run["id"],
                    category.category,
                    product["asin"],
                    product["title"],
                    product["brand"],
                    float(product["price"]) if product["price"] is not None else None,
                    product["currency"],
                    float(product["rating"]) if product["rating"] is not None else None,
                    product["rating_count"],
                    product["observed_at"],
                    json.dumps(product["specs"], ensure_ascii=False),
                ]
            )
        for asin, fields in evidence.items():
            tables["留存來源"].append(
                [run["id"], asin, json.dumps(fields, ensure_ascii=False, default=str)]
            )
    return tables


def build_report(db, settings, occurrence_id, task_id):
    row = occurrence(db, occurrence_id)
    if not row["results"]:
        raise ValueError("沒有成功批次可產報")
    tables = report_tables(db, occurrence_id)
    root = (
        Path(settings.scheduler_config_path).resolve().parent
        if settings.scheduler_config_path
        else Path.cwd()
    )
    destination = root / "exports" / "operations" / (task_id + ".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    edge = Side(style="thin", color="B8C5D1")
    for name, rows in tables.items():
        sheet = book.create_sheet(name)
        sheet.freeze_panes = "A2"
        for i, values in enumerate(rows, 1):
            for j, value in enumerate(values, 1):
                cell = sheet.cell(i, j, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(name="Arial", size=11, color="26384A")
                if i == 1:
                    cell.fill = PatternFill("solid", fgColor="18324D")
                    cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            sheet.row_dimensions[i].height = 45 if i == 1 else 60
        for j, label in enumerate(rows[0], 1):
            sheet.column_dimensions[get_column_letter(j)].width = (
                65 if label in {"內容", "商品名稱", "來源欄位 JSON"} else 24
            )
        sheet.auto_filter.ref = sheet.dimensions
    with destination.open("xb") as stream:
        destination.chmod(0o600)
        book.save(stream)
    loaded = load_workbook(destination, read_only=True, data_only=False)
    try:
        for name, rows in tables.items():
            expected = [tuple(None if v == "" else v for v in r) for r in rows]
            if list(loaded[name].values) != expected:
                raise ValueError("報表讀回不符")
            if any(c.data_type == "f" for r in loaded[name] for c in r):
                raise ValueError("不允許公式")
    finally:
        loaded.close()
    through = max(r["run_id"] for r in row["results"])
    sha = hashlib.sha256(destination.read_bytes()).hexdigest()
    record = {
        "schema_version": 1,
        "verification_status": "PASS",
        "source": row["spec"]["source"],
        "through_run_id": through,
        "sha256": sha,
        "workbook_schema_version": 5,
        "occurrence_id": occurrence_id,
        "run_ids": [r["run_id"] for r in row["results"]],
    }
    side = Path(str(destination) + ".verified.json")
    with side.open("x") as stream:
        side.chmod(0o600)
        json.dump(record, stream)
    with db.engine.begin() as conn:
        conn.execute(
            insert(operation_reports).values(
                task_id=task_id,
                source=row["spec"]["source"],
                through_run_id=through,
                path=str(destination),
                sha256=sha,
                created_at=utc_now(),
                workbook_version=5,
            )
        )
        from sqlalchemy import update

        conn.execute(update(tasks).where(tasks.c.id == task_id).values(run_id=through))
    return record


def enqueue_rebuild(db, settings, occurrence_id, task_id):
    from gummy.locking import process_lock
    from gummy.operations import config_binding

    row = occurrence(db, occurrence_id)
    if not row["results"] or row["status"] in {"QUEUED", "RUNNING"}:
        raise ValueError("請先等待採集完成，且至少有一個成功分類")
    with process_lock(db, "admission"), db.engine.begin() as conn:
        if conn.execute(select(tasks.c.id).where(tasks.c.id == task_id)).first():
            return task_id
        conn.execute(
            insert(tasks).values(
                id=task_id,
                kind="schedule-report",
                source=row["spec"]["source"],
                request={
                    "occurrence_id": occurrence_id,
                    "category": row["spec"]["name"],
                    "query": "已存批次重產",
                },
                binding=config_binding(settings),
                status="QUEUED",
                phase="queued",
                reserved_requests=0,
                cancel_requested=False,
                created_at=utc_now(),
                updated_at=utc_now(),
            )
        )
    return task_id
